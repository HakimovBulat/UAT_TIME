import logging
from telegram import (
    ReplyKeyboardMarkup,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
)
from telegram.ext import (
    Application,
    MessageHandler,
    filters,
    CommandHandler,
    ConversationHandler,
    CallbackQueryHandler,
)
from timetable import send_day_timetable, send_ring_time, CURRENT_WEEK_NUMBER
from openpyxl import load_workbook
import datetime
import json
import sqlite3


logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO
)
logger = logging.getLogger(__name__)
reply_keyboard = [["Сегодня", "Завтра"], ["Выбрать день", "Специальность"], ["Помощь", "Стоп"]]
markup = ReplyKeyboardMarkup(reply_keyboard, one_time_keyboard=False)
connection = sqlite3.connect("my_database.db")
cursor = connection.cursor()
connection.commit()
WEEK_NAMES = [
    "Понедельник",
    "Вторник",
    "Среда",
    "Четверг",
    "Пятница",
    "Суббота",
    "Воскресенье",
]
TOKEN = "8308113992:AAG32wAZzsTxpbmNlKn3Mf0SejY_1I_0wFY"
# TOKEN = "7521976097:AAHy6d-dRYbM0xB4KkNT5fiTQ7juhUe0NGI"
WEEK_NUMBER = 0

def delivery_report(err, msg):
    if err is not None:
        print(f'Ошибка доставки сообщения: {err}')
    else:
        print(f'Сообщение доставлено в {msg.topic()} [{msg.partition()}]')


async def today(update, context):
    day = WEEK_NAMES[datetime.datetime.today().weekday()]
    faculty_name, group_name = cursor.execute(
            """SELECT faculty, group_name FROM User WHERE id = ?""",
            (update.effective_user.id,),
        ).fetchone()
    if datetime.datetime.today().weekday() != 6:
        
        lessons_str = "\n".join(send_day_timetable(group_name, faculty_name, day))
        message = f"{day}, {group_name}: \n{lessons_str}"
    else:
        message = f"{day}, {group_name}: \nПар нет, так что можно отдохнуть"
    await update.message.reply_text(message)
    

async def select_week(update, context):
    await update.message.reply_text(
        f"Напишите номер нужной недели для просмотра расписания (текущая - {CURRENT_WEEK_NUMBER})"
    )
    return 1


async def select_day(update, context):
    global WEEK_NUMBER
    WEEK_NUMBER = int(update.message.text)
    keyboard = [
        [
            InlineKeyboardButton("Понедельник", callback_data="Понедельник"),
            InlineKeyboardButton("Вторник", callback_data="Вторник"),
        ],
        [
            InlineKeyboardButton("Среда", callback_data="Среда"),
            InlineKeyboardButton("Четверг", callback_data="Четверг"),
        ],
        [
            InlineKeyboardButton("Пятница", callback_data="Пятница"),
            InlineKeyboardButton("Суббота", callback_data="Суббота"),
        ],
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    await update.message.reply_text("Выберите день:", reply_markup=reply_markup)
    return 2


async def stop(update, context):
    await update.message.reply_text("Пожалуйста введи нужную команду заново")
    return ConversationHandler.END


async def button_day(update, context) -> None:
    query = update.callback_query
    await query.answer()
    day = query.data
    faculty_name, group_name = cursor.execute(
        """SELECT faculty, group_name FROM User WHERE id = ?""",
        (update.effective_user.id,),
    ).fetchone()
    lessons_str = "\n".join(
        send_day_timetable(group_name, faculty_name, day, WEEK_NUMBER)
    )
    message = f"{day}, {group_name}: \n{lessons_str}"
    await query.edit_message_text(text=message)
    return ConversationHandler.END


async def tomorrow(update, context):
    day = WEEK_NAMES[(datetime.datetime.today().weekday() + 1) % 7]
    faculty_name, group_name = cursor.execute(
            """SELECT faculty, group_name FROM User WHERE id = ?""",
            (update.effective_user.id,),
        ).fetchone()
    if datetime.datetime.today().weekday() != 5:
        
        if day == 0:
            lessons_str = "\n".join(send_day_timetable(group_name, faculty_name, day, datetime.datetime.now().isocalendar()[1] - 34))
        else:
            lessons_str = "\n".join(send_day_timetable(group_name, faculty_name, day))
        message = f"{day}, {group_name}: \n{lessons_str}"
    else:
        message = f"{day}, {group_name}: \nПар нет"
    await update.message.reply_text(message)


async def echo(update, context):
    if update.message.text == "Сегодня":
        await today(update, context)
    elif update.message.text == "Завтра":
        await tomorrow(update, context)
    elif update.message.text == "Специальность":
        await select_faculty(update, context)
    elif update.message.text == "Выбрать день":
        await select_day(update, context)

async def ring(update, context):
    message = f"Звонок в {send_ring_time()}"
    if message == "Звонок в Пары закончились":
        message = "Пары закончились"
    await update.message.reply_text(message)


async def help_command(update, context):
    await update.message.reply_text(
        "Это бот для просмотра расписаний. Вот мои команды:  \
            \n/stop - Нажать, если программа перестанет работать \
            \n/day - Выбрать конкретные неделю и день\
            \n/today - Расписание на сегодня \
            \n/tomorrow - Расписание на завтра \
            \n/faculty - Выбрать специальность \
            \n/group - Выбрать группу по специальности",
    )


async def start(update, context):
    user = update.effective_user
    cursor = connection.cursor()
    cursor.execute(
        """ CREATE TABLE IF NOT EXISTS User (id INTEGER PRIMARY KEY, group_name VARCHAR(50), faculty VARCHAR(50))"""
    )
    connection.commit()
    if not cursor.execute(
        """SELECT 1 FROM User WHERE id = ? """, (user.id,)
    ).fetchone():
        cursor.execute(
            """INSERT INTO User (id, group_name, faculty) VALUES (?, ?, ?)""",
            (user.id, "ИСП(п)4122", "ИСП ПР"),
        )
        connection.commit()
    await update.message.reply_html(
        rf"Привет, {user.mention_html()}! Жми /faculty для выбора специальности или /help для просмотра команд",
        reply_markup=markup,
    )


async def select_faculty(update, context):
    faculties = load_workbook("timetable.xlsx").sheetnames
    keyboard = []
    for faculty in faculties:
        keyboard.append([InlineKeyboardButton(faculty, callback_data=faculty)])
    reply_markup = InlineKeyboardMarkup(keyboard)
    await update.message.reply_html(
        rf"Выбери специальность: ", reply_markup=reply_markup
    )
    print(9999999999999999)
    await button_faculty(update, context)
    return 2


async def button_faculty(update, context):
    print(8888888888888888)
    query = update.callback_query
    await query.answer()
    cursor = connection.cursor()
    cursor.execute(
        """UPDATE User SET faculty = ? WHERE id = ?""",
        (query.data, update.effective_user.id),
    )
    connection.commit()
    await query.edit_message_text(
        text=f"Отлично, выша специльность: {query.data}. Жми /group"
    )
    return ConversationHandler.END


async def select_group(update, context):
    keyboard = []
    print(9999999999999999)
    faculty_name = cursor.execute(
        """SELECT faculty FROM User WHERE id = ?""", (update.effective_user.id,)
    ).fetchone()[0]
    print(9999999999999999)

    for group in load_workbook("timetable.xlsx")[faculty_name]["9"]:
        print(group.value)
        if group.value is not None and group.value not in [
            "День недели",
            "Время",
            "№ пары",
            "zz",

        ]:
            
            keyboard.append(
                [InlineKeyboardButton(group.value, callback_data=group.value)]
            )
    reply_markup = InlineKeyboardMarkup(keyboard)
    await update.message.reply_text(f"Выберите группу: ", reply_markup=reply_markup)
    return 2


async def button_group(update, context) -> None:
    query = update.callback_query
    await query.answer()
    cursor.execute(
        """UPDATE User SET group_name = ? WHERE id = ?""",
        (query.data, update.effective_user.id),
    )
    connection.commit()
    await query.edit_message_text(text="Можете посмотреть расписание на сегодня - /today или на завтра - /tomorrow")
    return ConversationHandler.END


def main():
    application = Application.builder().token(TOKEN).build()
    application.add_handler(CommandHandler("start", start))
    select_faculty_handler = ConversationHandler(
        entry_points=[CommandHandler("faculty", select_faculty)],
        states={
            2: [CallbackQueryHandler(button_faculty)],
        },
        fallbacks=[CommandHandler("stop", stop)],
    )
    application.add_handler(select_faculty_handler)

    select_group_handler = ConversationHandler(
        entry_points=[CommandHandler("group", select_group)],
        states={
            2: [CallbackQueryHandler(button_group)],
        },
        fallbacks=[CommandHandler("stop", stop)],
    )
    application.add_handler(select_group_handler)

    application.add_handler(CommandHandler("help", help_command))
    application.add_handler(CommandHandler("ring", ring))
    application.add_handler(CommandHandler("today", today))
    application.add_handler(CommandHandler("tomorrow", tomorrow))
    select_day_handler = ConversationHandler(
        entry_points=[CommandHandler("day", select_week)],
        states={
            1: [MessageHandler(filters.TEXT & ~filters.COMMAND, select_day)],
            2: [CallbackQueryHandler(button_day)],
        },
        fallbacks=[CommandHandler("stop", stop)],
    )
    application.add_handler(select_day_handler)
    text_handler = MessageHandler(filters.TEXT & ~filters.COMMAND, echo)
    application.add_handler(text_handler)
    application.run_polling()


if __name__ == "__main__":
    main()