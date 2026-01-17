from openpyxl import load_workbook
import datetime
print(datetime.datetime.now().isocalendar()[1] - 35)
wb = load_workbook("7 семестр Расписание 4 курса.xlsx")
for group in load_workbook("7 семестр Расписание 4 курса.xlsx")["ИСП ПР"]["9"]:
    print(group.value)